#!/usr/bin/env python3
"""Extract text from source.docx and split into sections for translation."""
import json
import os
from docx import Document

PROJECT_DIR = os.path.dirname(os.path.abspath(__file__))
INPUT_FILE = os.path.join(PROJECT_DIR, "source.docx")
OUTPUT_DIR = os.path.join(PROJECT_DIR, "sections")

os.makedirs(OUTPUT_DIR, exist_ok=True)

doc = Document(INPUT_FILE)

# Extract all paragraphs with their styles
paragraphs = []
for i, para in enumerate(doc.paragraphs):
    paragraphs.append({
        "index": i,
        "text": para.text,
        "style": para.style.name if para.style else "Normal",
    })

# Extract table cells
table_cells = []
for ti, table in enumerate(doc.tables):
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            key = f"t_{ti}_r{ri}_c{ci}"
            table_cells.append({
                "key": key,
                "text": cell.text,
            })

# Print summary
total_chars = sum(len(p["text"]) for p in paragraphs)
non_empty = [p for p in paragraphs if p["text"].strip()]
print(f"Total paragraphs: {len(paragraphs)}")
print(f"Non-empty paragraphs: {len(non_empty)}")
print(f"Total characters: {total_chars}")
print(f"Table cells: {len(table_cells)}")

# Split into sections of roughly equal size (aim for ~3000 chars per section)
SECTION_SIZE = 3000
sections = []
current_section = []
current_chars = 0

for p in paragraphs:
    current_section.append(p)
    current_chars += len(p["text"])
    if current_chars >= SECTION_SIZE:
        sections.append(current_section)
        current_section = []
        current_chars = 0

if current_section:
    sections.append(current_section)

print(f"Number of sections: {len(sections)}")

# Save each section
for i, section in enumerate(sections):
    section_file = os.path.join(OUTPUT_DIR, f"section_{i:02d}.json")
    with open(section_file, "w", encoding="utf-8") as f:
        json.dump(section, f, ensure_ascii=False, indent=2)
    chars = sum(len(p["text"]) for p in section)
    paras = len(section)
    print(f"  Section {i}: {paras} paragraphs, {chars} chars")

# Save table cells separately
table_file = os.path.join(OUTPUT_DIR, "table_cells.json")
with open(table_file, "w", encoding="utf-8") as f:
    json.dump(table_cells, f, ensure_ascii=False, indent=2)
print(f"Table cells saved to: {table_file}")

# Save metadata
with open(os.path.join(OUTPUT_DIR, "metadata.json"), "w", encoding="utf-8") as f:
    json.dump({
        "total_paragraphs": len(paragraphs),
        "total_sections": len(sections),
        "total_table_cells": len(table_cells),
        "input_file": INPUT_FILE,
    }, f, ensure_ascii=False, indent=2)

# Save document structure info (styles, formatting)
structure = []
for para in doc.paragraphs:
    runs_info = []
    for run in para.runs:
        runs_info.append({
            "text": run.text,
            "bold": run.bold,
            "italic": run.italic,
            "underline": run.underline,
            "font_name": run.font.name,
            "font_size": str(run.font.size) if run.font.size else None,
        })
    structure.append({
        "style": para.style.name if para.style else "Normal",
        "alignment": str(para.alignment) if para.alignment else None,
        "runs": runs_info,
    })

with open(os.path.join(OUTPUT_DIR, "structure.json"), "w", encoding="utf-8") as f:
    json.dump(structure, f, ensure_ascii=False, indent=2)

print("\nFirst few paragraphs:")
for p in paragraphs[:10]:
    if p["text"].strip():
        print(f"  [{p['style']}] {p['text'][:100]}")

print("\nTable cells:")
for tc in table_cells:
    print(f"  [{tc['key']}] {tc['text'][:80]}...")
