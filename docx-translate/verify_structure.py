#!/usr/bin/env python3
"""Verify translated DOCX structure matches source — sections, clauses, completeness."""
import json
import os
import re
import sys
from docx import Document

PROJECT_DIR = os.path.dirname(os.path.abspath(__file__))
SOURCE_FILE = os.path.join(PROJECT_DIR, "source.docx")

# Accept output file as argument or find latest variant
if len(sys.argv) > 1:
    OUTPUT_FILE = sys.argv[1]
else:
    import glob
    variants = sorted(glob.glob(os.path.join(PROJECT_DIR, "variant *.docx")))
    if not variants:
        print("ERROR: No variant files found")
        sys.exit(1)
    OUTPUT_FILE = variants[-1]

print(f"Source:  {os.path.basename(SOURCE_FILE)}")
print(f"Output:  {os.path.basename(OUTPUT_FILE)}")
print("=" * 60)

src = Document(SOURCE_FILE)
out = Document(OUTPUT_FILE)

errors = []
warnings = []

# 1. Paragraph count
src_count = len(src.paragraphs)
out_count = len(out.paragraphs)
if src_count != out_count:
    errors.append(f"Paragraph count mismatch: source={src_count}, output={out_count}")
else:
    print(f"[OK] Paragraph count: {src_count}")

# 2. Table count
src_tables = len(src.tables)
out_tables = len(out.tables)
if src_tables != out_tables:
    errors.append(f"Table count mismatch: source={src_tables}, output={out_tables}")
else:
    print(f"[OK] Table count: {src_tables}")

# 3. Table cell count
src_cells = sum(len(row.cells) for table in src.tables for row in table.rows)
out_cells = sum(len(row.cells) for table in out.tables for row in table.rows)
if src_cells != out_cells:
    errors.append(f"Table cell count mismatch: source={src_cells}, output={out_cells}")
else:
    print(f"[OK] Table cell count: {src_cells}")

# 4. Heading styles preserved
for i, (sp, op) in enumerate(zip(src.paragraphs, out.paragraphs)):
    src_style = sp.style.name if sp.style else "Normal"
    out_style = op.style.name if op.style else "Normal"
    if src_style != out_style:
        errors.append(f"Style mismatch at para {i}: source='{src_style}', output='{out_style}'")

style_mismatches = [e for e in errors if "Style mismatch" in e]
if not style_mismatches:
    print(f"[OK] All heading styles preserved")
else:
    for e in style_mismatches:
        print(f"[FAIL] {e}")

# 5. Clause number extraction and comparison
clause_pattern = re.compile(r'^(\d+\.\d+\.?\d*\.?)')

src_clauses = []
out_clauses = []
for i, (sp, op) in enumerate(zip(src.paragraphs, out.paragraphs)):
    src_match = clause_pattern.match(sp.text.strip())
    out_match = clause_pattern.match(op.text.strip())
    if src_match:
        src_clauses.append((i, src_match.group(1)))
    if out_match:
        out_clauses.append((i, out_match.group(1)))

src_clause_nums = set(c[1] for c in src_clauses)
out_clause_nums = set(c[1] for c in out_clauses)

missing = src_clause_nums - out_clause_nums
extra = out_clause_nums - src_clause_nums

if not missing and not extra:
    print(f"[OK] All {len(src_clause_nums)} clause numbers present")
else:
    if missing:
        errors.append(f"Missing clauses in output: {sorted(missing)}")
    if extra:
        warnings.append(f"Extra clauses in output: {sorted(extra)}")

# 6. Section headings (numbered sections 1-9)
section_pattern = re.compile(r'^(\d+)\.\s+\S')
src_sections = []
out_sections = []
for i, (sp, op) in enumerate(zip(src.paragraphs, out.paragraphs)):
    sm = section_pattern.match(sp.text.strip())
    om = section_pattern.match(op.text.strip())
    if sm:
        src_sections.append((i, sm.group(1), sp.text.strip()[:60]))
    if om:
        out_sections.append((i, om.group(1), op.text.strip()[:60]))

src_sec_nums = [s[1] for s in src_sections]
out_sec_nums = [s[1] for s in out_sections]

if src_sec_nums == out_sec_nums:
    print(f"[OK] All {len(src_sec_nums)} section headings present")
else:
    missing_secs = set(src_sec_nums) - set(out_sec_nums)
    if missing_secs:
        errors.append(f"Missing section headings: {sorted(missing_secs)}")

# 7. Content completeness — no dropped translations
dropped = []
for i, (sp, op) in enumerate(zip(src.paragraphs, out.paragraphs)):
    if sp.text.strip() and not op.text.strip():
        dropped.append((i, sp.text[:60]))

if not dropped:
    print(f"[OK] No dropped translations (all non-empty source paras have output)")
else:
    for idx, txt in dropped:
        errors.append(f"Dropped translation at para {idx}: '{txt}...'")

# 8. Table cell content check
for ti, (st, ot) in enumerate(zip(src.tables, out.tables)):
    for ri, (sr, orow) in enumerate(zip(st.rows, ot.rows)):
        for ci, (sc, oc) in enumerate(zip(sr.cells, orow.cells)):
            if sc.text.strip() and not oc.text.strip():
                errors.append(f"Empty table cell t_{ti}_r{ri}_c{ci} (source has content)")
            elif sc.text.strip() and sc.text == oc.text:
                warnings.append(f"Table cell t_{ti}_r{ri}_c{ci} appears untranslated")

# 9. Terminology consistency checks
full_text = "\n".join(p.text for p in out.paragraphs)
full_text += "\n" + "\n".join(
    c.text for t in out.tables for r in t.rows for c in r.cells
)

forbidden = {
    "თანამშრომელი": "should be დასაქმებული (Labor Code Art. 3)",
    "მუშაკი": "should be დასაქმებული (Labor Code Art. 3)",
    "ლოკალური ნორმატიული აქტები": "Russicism — should be შინაგანაწესი (Labor Code Art. 23)",
    "ადგილობრივი ნორმატიული": "Russicism — should be შინაგანაწესი (Labor Code Art. 23)",
    "თანამდებობრივი ინსტრუქცია": "Russicism — should be სამუშაოს აღწერილობა (Labor Code Art. 14)",
    "არასრული დასაქმება": "should be არასრული სამუშაო განაკვეთი (Labor Code Art. 16)",
}
for term, fix in forbidden.items():
    if term in full_text:
        errors.append(f"Forbidden term found: '{term}' ({fix})")

russian_months = ["января", "февраля", "марта", "апреля", "мая", "июня",
                   "июля", "августа", "сентября", "октября", "ноября", "декабря"]
for month in russian_months:
    if month in full_text:
        errors.append(f"Untranslated Russian month found: '{month}'")

required = ["დასაქმებული", "დამსაქმებელი", "შინაგანაწესი", "სამუშაოს აღწერილობა"]
for term in required:
    if term not in full_text:
        warnings.append(f"Expected term not found: '{term}'")

# Summary
print("=" * 60)
if not errors and not warnings:
    print("ALL CHECKS PASSED")
else:
    if warnings:
        print(f"\nWARNINGS ({len(warnings)}):")
        for w in warnings:
            print(f"  [WARN] {w}")
    if errors:
        print(f"\nERRORS ({len(errors)}):")
        for e in errors:
            print(f"  [FAIL] {e}")
        sys.exit(1)
    else:
        print("\nAll checks passed (with warnings)")
