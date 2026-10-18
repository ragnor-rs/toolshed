#!/usr/bin/env python3
"""Generate a synthetic Russian labor contract as ../source.docx, to try the pipeline on.

Everyone and everything in it is invented. It is deliberately short — 30-odd paragraphs
rather than the ~98 of a real contract — so a full extract/translate/stitch/verify cycle
is cheap to run. It carries the structural features the pipeline actually has to survive:
Heading 1 and Heading 2 styles, numbered clauses at two depths, a dash-prefixed list, a
multi-line paragraph using line breaks, and a two-cell signature table.
"""
import os

from docx import Document

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), os.pardir, "source.docx")

SECTIONS = [
    ("1. ПРЕДМЕТ ДОГОВОРА", [
        "1.1. Работодатель принимает Работника на должность бариста-кассира.",
        "1.2. Работа по настоящему договору является для Работника основной.",
        "1.3. Договор заключается на определённый срок.",
    ]),
    ("2. ДОЛЖНОСТНЫЕ ОБЯЗАННОСТИ", [
        "2.1. Работник обязан знать и соблюдать следующее:",
        "- устройство и настройку кофе-машины;",
        "- рецептуру эспрессо, капучино, латте, американо;",
        "- правила кассовой дисциплины;",
        "2.2. Работник обслуживает посетителей за барной стойкой.",
        "2.3. Работник соблюдает должностную инструкцию.",
    ]),
    ("3. РАБОЧЕЕ ВРЕМЯ", [
        "3.1. Работнику устанавливается неполное рабочее время — 0,5 ставки.",
        "3.2. Выходные дни предоставляются согласно внутреннему распорядку.",
    ]),
    ("4. ОПЛАТА ТРУДА", [
        "4.1. Заработная плата выплачивается два раза в месяц.",
        "4.2. Работнику предоставляется ежегодный оплачиваемый отпуск 24 рабочих дня.",
    ]),
    ("5. ПРЕКРАЩЕНИЕ ТРУДОВОГО ДОГОВОРА", [
        "5.1. Договор может быть прекращён по соглашению сторон.",
        "5.1.1. Работник предупреждает Работодателя за две недели.",
        "5.1.2. Работодатель знакомит Работника с приказом под подпись.",
        "5.2. Договор составлен в двух экземплярах, имеющих равную юридическую силу.",
    ]),
]


def main():
    doc = Document()

    doc.add_paragraph("ТРУДОВОЙ ДОГОВОР № _001_", style="Heading 1")
    doc.add_paragraph("город Тбилиси, Грузия\n «_01_»__марта___ 2026 года")
    doc.add_paragraph(
        "С одной стороны, Индивидуальный предприниматель _Иванова Анна_____\n"
        "Идентификационный номер: __000000000_____________________________\n"
        "Адрес регистрации: __Грузия, Тбилиси, ул. Примерная, д. 1________"
    )
    doc.add_paragraph("далее именуемый «Работодатель» , а с другой стороны,")
    doc.add_paragraph(
        "физическое лицо____Петров  Иван________________________________\n"
        "Паспорт/ID: ___000000000_______________________________________\n"
        "Адрес: ___Тбилиси, ул. Примерная, д. 2, кв. 2__________________"
    )
    doc.add_paragraph("далее именуемый «Работник», заключили настоящий договор.")

    for heading, clauses in SECTIONS:
        doc.add_paragraph(heading, style="Heading 2")
        for clause in clauses:
            doc.add_paragraph(clause)
        doc.add_paragraph("")

    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = (
        "Работодатель\nИндивидуальный предприниматель\nИванова Анна\n"
        "ID 000000000\nмоб. номер +995 000 00 00 00\nПодпись ____________\nДата ____________"
    )
    table.rows[0].cells[1].text = (
        "Работник\nПетров Иван\nномер паспорта 000000000\n"
        "моб. номер +995 000 00 00 00\nПодпись ____________\nДата ____________"
    )

    doc.save(OUT)
    paras = len(doc.paragraphs)
    cells = sum(len(r.cells) for t in doc.tables for r in t.rows)
    print(f"Wrote {os.path.realpath(OUT)}")
    print(f"  {paras} paragraphs, {len(doc.tables)} table, {cells} cells")


if __name__ == "__main__":
    main()
